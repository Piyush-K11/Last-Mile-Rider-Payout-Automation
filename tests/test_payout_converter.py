"""
Tests for the Rider Payout Converter.

Run from the project root (no extra dependencies needed):

    python -m unittest discover -s tests -v

Numbers are read from COMPANY_CONFIG rather than hard-coded, so editing the
sample rate card doesn't break the tests.
"""
import contextlib
import io
import os
import sys
import tempfile
import time
import unittest
import zipfile
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import app as web  # noqa: E402
from payout_engine import (COMPANY_CONFIG, MissingColumnError, build, login_amount,  # noqa: E402
                           preview_summary, store_has_mg)

SAMPLE = os.path.join(ROOT, 'sample_data')
SHIPMENTS = os.path.join(SAMPLE, 'shipments_week.xlsx')
EXAMPLE_IN = os.path.join(SAMPLE, 'example', 'example_input__ALPHA_EAST_HUB.xlsx')
EXAMPLE_OUT = os.path.join(SAMPLE, 'example', 'example_output__ALPHA_EAST_HUB.xlsx')
EX_VENDOR, EX_STORE = 'ALPHA', 'EAST HUB'


def sample_path(vendor, store):
    return os.path.join(SAMPLE, 'raw', f'{vendor}__{store}'.replace(' ', '_') + '.xlsx')


def all_stores():
    return [(v, s) for v, stores in COMPANY_CONFIG.items() for s in stores]


def expected_final_payout(raw_file, cfg):
    """Independent per-rider calculation straight from the raw trips (pandas only)."""
    df = pd.read_excel(raw_file, sheet_name='Sheet1')
    g = df.groupby('RIDER NAME').agg(days=('DATE', 'nunique'), orders=('TOTAL DELIVERED ORDERS', 'sum'))
    g['incentive'] = [login_amount(int(d), int(o), cfg['tiers']) for d, o in zip(g.days, g.orders)]
    g['payout'] = g.orders * cfg['rate']
    return g


class SampleDataTests(unittest.TestCase):
    def test_every_configured_store_has_a_sample_file(self):
        for vendor, store in all_stores():
            with self.subTest(vendor=vendor, store=store):
                self.assertTrue(os.path.exists(sample_path(vendor, store)))

    def test_every_sample_file_previews_and_builds(self):
        """All vendor/store payout models run end to end, with the shipment check where it applies."""
        with tempfile.TemporaryDirectory() as tmp:
            for vendor, store in all_stores():
                with self.subTest(vendor=vendor, store=store):
                    raw = sample_path(vendor, store)
                    summary = preview_summary(raw, company=vendor, store=store)
                    self.assertGreater(summary['rider_count'], 0)
                    out = os.path.join(tmp, 'out.xlsx')
                    shipment = None if store_has_mg(vendor, store) else SHIPMENTS
                    build(raw, out, company=vendor, store=store, shipment_path=shipment)
                    self.assertTrue(load_workbook(out).sheetnames)


class EngineTests(unittest.TestCase):
    def test_final_payout_matches_independent_calculation(self):
        cfg = COMPANY_CONFIG[EX_VENDOR][EX_STORE]
        expected = expected_final_payout(EXAMPLE_IN, cfg)
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, 'out.xlsx')
            build(EXAMPLE_IN, out, company=EX_VENDOR, store=EX_STORE)
            ws = load_workbook(out)['FINAL PAYOUT']
            rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] != 'TOTAL']
        self.assertEqual({r[0] for r in rows}, set(expected.index))
        for name, days, orders, incentive, payout, _total, insurance, _final, _dcpo in rows:
            e = expected.loc[name]
            self.assertEqual((days, orders, incentive, payout, insurance),
                             (e.days, e.orders, e.incentive, e.payout, cfg['insurance']), name)

    def test_preview_totals_agree_with_independent_calculation(self):
        cfg = COMPANY_CONFIG[EX_VENDOR][EX_STORE]
        e = expected_final_payout(EXAMPLE_IN, cfg)
        summary = preview_summary(EXAMPLE_IN, company=EX_VENDOR, store=EX_STORE)
        riders = len(e)
        self.assertEqual(summary['rider_count'], riders)
        self.assertEqual(summary['totals']['delivered'], e.orders.sum())
        self.assertEqual(summary['totals']['final_payout'],
                         e.payout.sum() + e.incentive.sum() - cfg['insurance'] * riders)

    def test_committed_example_output_matches_its_input(self):
        """The example workbook shipped in the repo must agree with the numbers recomputed from its input."""
        cfg = COMPANY_CONFIG[EX_VENDOR][EX_STORE]
        expected = expected_final_payout(EXAMPLE_IN, cfg)
        ws = load_workbook(EXAMPLE_OUT)['FINAL PAYOUT']
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] == 'TOTAL':
                continue
            self.assertEqual((r[1], r[2], r[3], r[4]),
                             tuple(expected.loc[r[0], ['days', 'orders', 'incentive', 'payout']]), r[0])

    def test_incentive_tiers(self):
        tiers = COMPANY_CONFIG[EX_VENDOR][EX_STORE]['tiers']
        days, orders, amount = tiers[0]
        self.assertEqual(login_amount(days, orders, tiers), amount)
        self.assertLess(login_amount(days, orders - 1, tiers), amount)      # one order short of the top tier
        lowest_days = min(t[0] for t in tiers)
        self.assertEqual(login_amount(lowest_days - 1, 10 ** 6, tiers), 0)  # too few days, however many orders

    def test_mixed_date_formats_count_each_day_once(self):
        """A real Excel date and a text date for the same calendar day must not count as two days."""
        rows = [
            {'DATE': datetime(2026, 9, 7), 'RIDER NAME': 'Test Rider', 'TOTAL DELIVERED ORDERS': 10},
            {'DATE': '08/09/2026', 'RIDER NAME': 'Test Rider', 'TOTAL DELIVERED ORDERS': 10},
            {'DATE': datetime(2026, 9, 8), 'RIDER NAME': 'Test Rider', 'TOTAL DELIVERED ORDERS': 5},
            {'DATE': '13/09/2026', 'RIDER NAME': 'Test Rider', 'TOTAL DELIVERED ORDERS': 10},
        ]
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'mixed.xlsx')
            pd.DataFrame(rows).to_excel(path, index=False)
            summary = preview_summary(path, company=EX_VENDOR, store=EX_STORE)
        self.assertEqual(summary['riders'][0]['days'], 3)  # Sep 7, 8, 13

    def test_missing_columns_are_reported_precisely(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, 'bad.xlsx')
            pd.DataFrame({'foo': [1]}).to_excel(path, index=False)
            with self.assertRaises(MissingColumnError) as ctx:
                preview_summary(path, company=EX_VENDOR, store=EX_STORE)
        self.assertIn('DATE', ctx.exception.missing)
        self.assertEqual(ctx.exception.found, ['foo'])


class WebAppTests(unittest.TestCase):
    def setUp(self):
        # keep every upload/output in a throwaway folder, never in the project directory
        self.tmp = tempfile.TemporaryDirectory()
        self._saved = (web.UPLOAD_DIR, web.OUTPUT_DIR)
        web.UPLOAD_DIR = os.path.join(self.tmp.name, 'uploads')
        web.OUTPUT_DIR = os.path.join(self.tmp.name, 'outputs')
        os.makedirs(web.UPLOAD_DIR)
        os.makedirs(web.OUTPUT_DIR)
        self.client = web.app.test_client()

    def tearDown(self):
        web.UPLOAD_DIR, web.OUTPUT_DIR = self._saved
        self.tmp.cleanup()

    @staticmethod
    def file(path, name=None):
        with open(path, 'rb') as fh:
            return io.BytesIO(fh.read()), name or os.path.basename(path)

    @staticmethod
    def settle(response):
        """Read a streamed response fully and release its file handle (keeps the test run warning-free)."""
        response.get_data()
        response.close()
        return response

    def get(self, url):
        return self.settle(self.client.get(url))

    def post(self, endpoint, **fields):
        with contextlib.redirect_stderr(io.StringIO()):  # the app logs tracebacks for handled errors
            return self.settle(self.client.post(endpoint, data=fields, content_type='multipart/form-data'))

    def test_pages_and_static_assets_load(self):
        for url in ('/', '/rate-card', '/static/logo.png'):
            self.assertEqual(self.get(url).status_code, 200, url)
        html = self.get('/').data.decode()
        for vendor in COMPANY_CONFIG:
            self.assertIn(vendor, html)

    def test_preview(self):
        r = self.post('/preview', company=EX_VENDOR, store=EX_STORE, raw_file=self.file(EXAMPLE_IN))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.get_json()['rider_count'], 14)

    def test_generate_returns_a_valid_workbook(self):
        r = self.post('/generate', company=EX_VENDOR, store=EX_STORE, raw_file=self.file(EXAMPLE_IN))
        self.assertEqual(r.status_code, 200)
        self.assertIn('attachment', r.headers['Content-Disposition'])
        self.assertEqual(load_workbook(io.BytesIO(r.data)).sheetnames, ['PAYOUT', 'PAYOUT DATA', 'FINAL PAYOUT'])

    def test_generate_with_shipment_file_adds_check_column(self):
        r = self.post('/generate', company=EX_VENDOR, store=EX_STORE, raw_file=self.file(EXAMPLE_IN),
                      shipment_file=self.file(SHIPMENTS))
        self.assertEqual(r.status_code, 200)
        header = [c.value for c in load_workbook(io.BytesIO(r.data))['PAYOUT'][1]]
        self.assertIn('SHIPMENT CHECK', header)

    def test_bulk_returns_zip_and_reports_failures(self):
        fields = {'row_count': '3'}
        for i, (v, s) in enumerate([(EX_VENDOR, 'NORTH HUB'), ('BETA', 'NORTH HUB')]):
            fields.update({f'company_{i}': v, f'store_{i}': s, f'raw_file_{i}': self.file(sample_path(v, s))})
        fields.update({'company_2': EX_VENDOR, 'store_2': 'EAST HUB'})  # row 3 has no file -> must fail gracefully
        r = self.post('/generate-bulk', **fields)
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers['X-Bulk-Summary'], '2 succeeded, 1 failed')
        names = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
        self.assertEqual(sum(n.endswith('.xlsx') for n in names), 2)
        self.assertIn('FAILURES.txt', names)

    def test_xls_and_other_extensions_are_rejected_with_a_clear_message(self):
        for name in ('legacy.xls', 'notes.txt', 'macro.xlsm'):
            r = self.post('/preview', company=EX_VENDOR, store=EX_STORE, raw_file=(io.BytesIO(b'x'), name))
            self.assertEqual(r.status_code, 400, name)
            self.assertIn('.xlsx', r.get_json()['error'])
        r = self.post('/generate', company=EX_VENDOR, store=EX_STORE, raw_file=self.file(EXAMPLE_IN),
                      shipment_file=(io.BytesIO(b'x'), 'ship.xls'))
        self.assertEqual(r.status_code, 400)

    def test_unknown_vendor_or_store_is_rejected_before_any_file_is_written(self):
        for company, store in (('NOPE', 'NORTH HUB'), (EX_VENDOR, 'NOWHERE'), ('../../x', '../y')):
            r = self.post('/generate', company=company, store=store, raw_file=self.file(EXAMPLE_IN))
            self.assertEqual(r.status_code, 400, (company, store))
        self.assertEqual(os.listdir(web.UPLOAD_DIR) + os.listdir(web.OUTPUT_DIR), [])

    def test_hostile_upload_filename_cannot_escape_the_upload_folder(self):
        marker = 'escape_probe_upload.xlsx'
        r = self.post('/preview', company=EX_VENDOR, store=EX_STORE,
                      raw_file=(self.file(EXAMPLE_IN)[0], f'../../{marker}'))
        self.assertEqual(r.status_code, 200)
        self.assertFalse(os.path.exists(os.path.join(os.path.dirname(self.tmp.name), marker)))
        self.assertTrue(any(f.endswith(marker) for f in os.listdir(web.UPLOAD_DIR)))

    def test_wrong_columns_give_an_actionable_error(self):
        bad = io.BytesIO()
        pd.DataFrame({'foo': [1]}).to_excel(bad, index=False)
        bad.seek(0)
        r = self.post('/preview', company=EX_VENDOR, store=EX_STORE, raw_file=(bad, 'bad.xlsx'))
        self.assertEqual(r.status_code, 400)
        body = r.get_json()
        self.assertIn('DATE', body['missing_columns'])
        self.assertEqual(body['found_columns'], ['foo'])

    def test_old_scratch_files_are_purged_and_fresh_ones_kept(self):
        old, fresh = (os.path.join(web.UPLOAD_DIR, n) for n in ('old.xlsx', 'fresh.xlsx'))
        for p in (old, fresh):
            open(p, 'w').close()
        two_days_ago = time.time() - 2 * 24 * 3600
        os.utime(old, (two_days_ago, two_days_ago))
        self.assertEqual(web.purge_old_files(), 1)
        self.assertFalse(os.path.exists(old))
        self.assertTrue(os.path.exists(fresh))


if __name__ == '__main__':
    unittest.main()
