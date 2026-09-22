"""
Generate fictional sample input files for the Payout Converter.

Writes one raw weekly attendance/trip workbook per configured vendor/store into
sample_data/raw/, one shared shipment file (sample_data/shipments_week.xlsx)
for the shipment cross-check feature, and a worked input -> output example
in sample_data/example/. Everything is randomly generated with a
fixed seed — no real riders, trips, vendors or amounts.

The generator reads COMPANY_CONFIG, so it stays in sync with whatever vendors
and stores you define there. Re-run it after changing the config:

    python sample_data/generate_sample_data.py
"""
import os
import random
import sys
from datetime import date, datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from payout_engine import COMPANY_CONFIG, build, login_amount  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
RAW_DIR = os.path.join(HERE, 'raw')
WEEK_START = date(2026, 9, 7)  # a Monday; the sample week runs Mon-Sun

FIRST_NAMES = [
    'Aarav', 'Vivaan', 'Aditya', 'Arjun', 'Kabir', 'Ishaan', 'Reyansh', 'Ayaan', 'Krish', 'Dev',
    'Neel', 'Yash', 'Karan', 'Manav', 'Nikhil', 'Pranav', 'Sahil', 'Tarun', 'Varun', 'Zayan',
    'Anaya', 'Diya', 'Ira', 'Kiara', 'Meera', 'Naina', 'Riya', 'Saanvi', 'Tara', 'Zoya',
]
LAST_NAMES = [
    'Mehta', 'Nair', 'Kapoor', 'Iyer', 'Bose', 'Chopra', 'Desai', 'Gill', 'Joshi', 'Khanna',
    'Malhotra', 'Pillai', 'Rao', 'Saxena', 'Trivedi', 'Ahuja', 'Bhatt', 'Dutta', 'Sethi', 'Menon',
]

_trip_counter = [10000]


def next_trip_id():
    _trip_counter[0] += 1
    return _trip_counter[0]


def model_of(cfg):
    """Classify a store config into the payout model the engine will apply."""
    if cfg.get('distance_mg'):
        return 'distance'
    if cfg.get('mg_per_trip'):
        return 'mg_per_trip'
    if cfg.get('flat_per_day'):
        return 'flat'
    if cfg.get('cancelled_rate'):
        return 'cancelled'
    return 'per_order'


def pick_riders(rng, n, exclude_first=()):
    names = set()
    while len(names) < n:
        f, l = rng.choice(FIRST_NAMES), rng.choice(LAST_NAMES)
        if f in exclude_first:
            continue
        names.add(f'{f} {l}')
    return sorted(names)


def make_rows(rng, riders, store_label, model, orders_range=(6, 17), absent_rate=0.05,
              no_trip_rate=0.0):
    """One row per rider-day (occasionally two trips in a day, occasionally an
    absence marked 'A'). Returns (headers, rows)."""
    headers = ['DATE', 'RIDER NAME', 'STORE NAME', 'TRIP ID', 'IN TIME', 'OUT TIME',
               'TOTAL DELIVERED ORDERS']
    if model == 'cancelled':
        headers.append('TOTAL CANCELLED ORDERS')
    if model == 'distance':
        headers.append('DISTANCE')

    rows = []
    for name in riders:
        # a rider's "profile" drives how many days they work and how busy they are
        profile = rng.choices(['full', 'regular', 'part'], weights=[4, 3, 3])[0]
        days = {'full': 7, 'regular': 6, 'part': rng.choice([3, 4, 5])}[profile]
        busy = {'full': 0.9, 'regular': 0.75, 'part': 0.5}[profile]
        worked = sorted(rng.sample(range(7), days))
        for d in worked:
            day = WEEK_START + timedelta(days=d)
            if rng.random() < absent_rate:
                rows.append({'DATE': day, 'RIDER NAME': name, 'STORE NAME': store_label,
                             'TRIP ID': None, 'IN TIME': 'A', 'OUT TIME': 'A',
                             'TOTAL DELIVERED ORDERS': 0, 'TOTAL CANCELLED ORDERS': 0,
                             'DISTANCE': 0})
                continue
            trips = 2 if rng.random() < 0.10 else 1
            for _ in range(trips):
                lo, hi = orders_range
                delivered = max(1, int(rng.randint(lo, hi) * busy / (1.4 if trips == 2 else 1)))
                start_h = rng.choice([8, 9, 10])
                trip_id = None if rng.random() < no_trip_rate else next_trip_id()
                rows.append({
                    'DATE': day, 'RIDER NAME': name, 'STORE NAME': store_label,
                    'TRIP ID': trip_id,
                    'IN TIME': f'{start_h:02d}:{rng.randint(0, 59):02d}',
                    'OUT TIME': f'{start_h + rng.randint(6, 9):02d}:{rng.randint(0, 59):02d}',
                    'TOTAL DELIVERED ORDERS': delivered,
                    'TOTAL CANCELLED ORDERS': rng.choice([0, 0, 0, 1, 2]),
                    'DISTANCE': round(rng.uniform(18, 62), 1),
                })
    return headers, rows


def write_sheet(ws, headers, rows):
    ws.append(headers)
    for r in rows:
        ws.append([
            datetime.combine(r[h], datetime.min.time()) if h == 'DATE' else r.get(h)
            for h in headers
        ])
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + i)].width = 22
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        row[0].number_format = 'dd-mm-yyyy'


def build_raw_workbook(rng, vendor, store, cfg):
    model = model_of(cfg)
    store_label = f"{vendor}_{store}".lower().replace(' ', '_')
    wb = Workbook()
    ws = wb.active
    all_rows = []

    if cfg.get('mg'):
        # Two raw sheets: per-order riders + minimum-guarantee riders.
        ws.title = 'PER ORDER DATA'
        po_riders = pick_riders(rng, 8)
        h, po_rows = make_rows(rng, po_riders, store_label, model)
        write_sheet(ws, h, po_rows)
        ws_mg = wb.create_sheet('MG ORDERS')
        mg_riders = pick_riders(rng, 3, exclude_first={n.split()[0] for n in po_riders})
        _, mg_rows = make_rows(rng, mg_riders, store_label, model, orders_range=(9, 22),
                               absent_rate=0.0)
        write_sheet(ws_mg, h, mg_rows)
        all_rows = po_rows  # only per-order trips take part in the shipment check
        return wb, all_rows

    riders = pick_riders(rng, 10 if model == 'per_order' else 8)
    ws.title = 'Sheet1'
    no_trip = 0.12 if vendor == 'TAU' else 0.0  # logged in but no trip that day
    headers, rows = make_rows(rng, riders, store_label, model, no_trip_rate=no_trip)

    # An 'Other Vendor' rider shares this store; the engine must exclude them.
    if cfg.get('restrict_vendor'):
        outsider = 'Rahul'  # trip sheets often carry shortened names
        _, extra = make_rows(rng, [outsider], store_label, model)
        rows += extra
    write_sheet(ws, headers, rows)

    if vendor == 'ALPHA':
        # ALPHA raw files carry a second, reference-only sheet (copied through as-is).
        ws2 = wb.create_sheet('Sheet2')
        ws2.append(['NOTE', 'VALUE'])
        ws2.append(['Reporting week starts', WEEK_START.isoformat()])
        ws2.append(['Prepared by', 'Sample Data Generator'])

    if cfg.get('restrict_vendor'):
        ws_att = wb.create_sheet('ATTENDANCE')
        ws_att.append(['Name', 'Vendor', 'Days Present'])
        for n in riders:
            ws_att.append([n, cfg['restrict_vendor'], sum(1 for r in rows if r['RIDER NAME'] == n)])
        ws_att.append(['Rahul Verma', 'Other Vendor', 5])

    if vendor == 'KAPPA':
        ws_att = wb.create_sheet('ATTENDANCE')
        ws_att.append(['Name', 'Store', 'Days Present'])
        for n in riders:
            ws_att.append([n, store, sum(1 for r in rows if r['RIDER NAME'] == n)])
        ws_bank = wb.create_sheet('AC DETAILS')
        ws_bank.append(['Rider Name', 'Bank', 'Account No', 'IFSC'])
        for i, n in enumerate(riders, start=1):
            # obviously-fake placeholders, not real account data
            ws_bank.append([n, 'Sample Bank', f'XXXXXX{i:04d}', f'SAMP000{i:04d}'])
        wb.create_sheet('CAL').append(['scratch'])  # junk tab the engine deliberately ignores

    all_rows = rows
    return wb, all_rows


def build_shipment_workbook(rng, trips):
    """trips: list of (trip_id, delivered). Mostly consistent with the raw files,
    with a few deliberate discrepancies so the red/green check is visible."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Summary'
    ws.append(['Shipment summary (sample data)'])
    ws.append([])
    ws.append(['Trip ID', 'Total'])
    stats = {'match': 0, 'over_reported': 0, 'missing': 0}
    for trip_id, delivered in trips:
        roll = rng.random()
        if roll < 0.03:
            stats['missing'] += 1          # trip absent from shipment data -> NOT FOUND
            continue
        if roll < 0.07:
            ws.append([trip_id, max(0, delivered - rng.randint(1, 3))])
            stats['over_reported'] += 1    # rider claims more than shipped -> red
            continue
        ws.append([trip_id, delivered + (rng.choice([0, 0, 0, 1]))])
        stats['match'] += 1
    ws.column_dimensions['A'].width = 14
    ws.column_dimensions['B'].width = 10
    return wb, stats


# ---------------------------------------------------------------------------
# Worked example: a realistic INPUT workbook and the OUTPUT the converter
# produces from it. The input mirrors the full export layout of a typical raw
# file (19 columns, plus a manual calculation sheet); the converter only needs
# a few of those columns and ignores the rest.
# ---------------------------------------------------------------------------
EXAMPLE_DIR = os.path.join(HERE, 'example')
EXAMPLE_VENDOR, EXAMPLE_STORE = 'ALPHA', 'EAST HUB'
EXAMPLE_HEADERS = [
    'DATE', 'TRIP ID', 'RIDER ID', 'RIDER NAME', 'STORE NAME', 'IN TIME', 'OUT TIME',
    'TOTAL WORKING HOURS', 'TOTAL ASSIGN ORDERS', 'TOTAL DELIVERED ORDERS',
    'TOTAL UNDELIVERED ORDERS', 'DELIVERY PERCENTAGE', 'ORDER REMARK',
    'OPENING METER READING', 'CLOSING METER READING', 'DISTANCE',
    'VERIFIED ORDER COUNT', 'ROUTES', 'REMARK',
]
EXAMPLE_WIDTHS = [11, 12, 11, 27, 17, 10, 12, 22, 22, 24, 26, 22, 30, 24, 24, 11, 22, 10, 10]
# (days worked, delivered orders for the week) — chosen to hit every incentive case:
# top tier, second tier, enough days but too few orders, too few days, single day.
EXAMPLE_PROFILES = [(7, 175), (7, 160), (7, 148), (7, 131), (7, 97), (7, 88), (6, 123),
                    (6, 104), (6, 92), (6, 72), (6, 41), (5, 64), (3, 33), (1, 13)]
EXAMPLE_REMARKS = ['1 ORDER REPLACED', '1 ORDER DELIVERED MANUALLY', '1 ITEM SHORT']

HEADER_FILL_Y = PatternFill('solid', fgColor='FFFF00')


def _spread(rng, total, days):
    """Split a weekly order total across `days` days (each >= 1, sums exactly)."""
    parts = [max(1, total // days + rng.randint(-3, 3)) for _ in range(days)]
    while sum(parts) != total:
        i = rng.randrange(days)
        if sum(parts) < total:
            parts[i] += 1
        elif parts[i] > 1:
            parts[i] -= 1
    return parts


def build_example_pair():
    rng = random.Random(7)  # separate stream: never disturbs the other sample files
    cfg = COMPANY_CONFIG[EXAMPLE_VENDOR][EXAMPLE_STORE]
    rate, tiers = cfg['rate'], cfg['tiers']

    names = set()
    while len(names) < len(EXAMPLE_PROFILES):
        names.add(f'{rng.choice(FIRST_NAMES)} {rng.choice(LAST_NAMES)}')
    names = sorted(names)
    rng.shuffle(names)

    trips = []  # (date, rider, delivered)
    for name, (days, total) in zip(names, EXAMPLE_PROFILES):
        for d, orders in zip(sorted(rng.sample(range(7), days)), _spread(rng, total, days)):
            day = datetime.combine(WEEK_START + timedelta(days=d), datetime.min.time())
            if orders >= 10 and rng.random() < 0.3:      # sometimes two trips in one day
                first = rng.randint(3, orders - 3)
                trips += [(day, name, first), (day, name, orders - first)]
            else:
                trips.append((day, name, orders))
    trips.sort(key=lambda x: (x[0], x[1]))

    wb = Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(EXAMPLE_HEADERS)
    trip_id = 41000
    for day, name, delivered in trips:
        trip_id += rng.randint(1, 4)
        undelivered = rng.randint(0, max(1, round(delivered * 0.6)))
        assigned = delivered + undelivered
        remark = rng.choice(EXAMPLE_REMARKS) if rng.random() < 0.06 else None
        # this export type fills time/hour/rider-id columns with 0 — kept as-is on purpose
        ws.append([day, trip_id, 0, name, EXAMPLE_STORE, 0, 0, 0, assigned, delivered, undelivered,
                   round(delivered / assigned, 4), remark, None, None, None, None, None, None])
    for c in ws[1]:
        c.font = Font(name='Arial', size=12, bold=True)
        c.fill = HEADER_FILL_Y
        c.alignment = Alignment(horizontal='center', vertical='center')
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name='Arial', size=10)
        row[0].number_format = 'mm-dd-yy'
        row[11].number_format = '0.00%'
    for i, w in enumerate(EXAMPLE_WIDTHS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    # Sheet2: the manual per-rider calculation kept alongside the raw export.
    # The converter copies it through untouched; it is not used in any calculation.
    ws2 = wb.create_sheet('Sheet2')
    ws2.append(['Rider Name', 'Delivered Order', 'Login Days', 'Slab', 'Login Amount',
                'Per order earning', 'Earning'])
    for i, name in enumerate(sorted(names, key=str.lower), start=2):
        mine = [t for t in trips if t[1] == name]
        orders, days = sum(t[2] for t in mine), len({t[0] for t in mine})
        ws2.append([name, orders, days, rate, login_amount(days, orders, tiers),
                    f'=+D{i}*B{i}', f'=+E{i}+F{i}'])
    last = len(names) + 1
    ws2.append(['Total', f'=SUM(B2:B{last})', None, None, None, None, f'=SUM(G2:G{last})'])
    ws2.append([])
    ws2.append(['DCPO', f'=G{last + 1}/B{last + 1}'])
    for c in ws2[1]:
        c.font = Font(name='Arial', size=12, bold=True)
        c.fill = HEADER_FILL_Y
        c.alignment = Alignment(horizontal='center', vertical='center')
    for col, w in zip('ABCDEFG', [32.7, 14.9, 10.7, 13.6, 13.3, 16.6, 10]):
        ws2.column_dimensions[col].width = w

    os.makedirs(EXAMPLE_DIR, exist_ok=True)
    raw_path = os.path.join(EXAMPLE_DIR, f'example_input__{EXAMPLE_VENDOR}_{EXAMPLE_STORE}.xlsx'.replace(' ', '_'))
    final_path = os.path.join(EXAMPLE_DIR, f'example_output__{EXAMPLE_VENDOR}_{EXAMPLE_STORE}.xlsx'.replace(' ', '_'))
    wb.save(raw_path)
    build(raw_path, final_path, company=EXAMPLE_VENDOR, store=EXAMPLE_STORE)
    return raw_path, final_path, len(trips)


def main():
    rng = random.Random(2026)
    os.makedirs(RAW_DIR, exist_ok=True)
    for f in os.listdir(RAW_DIR):
        if f.endswith('.xlsx'):
            os.remove(os.path.join(RAW_DIR, f))

    shipment_trips = []
    count = 0
    for vendor, stores in COMPANY_CONFIG.items():
        for store, cfg in stores.items():
            wb, rows = build_raw_workbook(rng, vendor, store, cfg)
            fname = f"{vendor}__{store}".replace(' ', '_') + '.xlsx'
            wb.save(os.path.join(RAW_DIR, fname))
            count += 1
            has_mg = bool(cfg.get('mg') or cfg.get('mg_per_trip')
                          or cfg.get('distance_mg') or cfg.get('flat_per_day'))
            if not has_mg:  # the shipment check is skipped for MG/flat stores
                shipment_trips += [(r['TRIP ID'], r['TOTAL DELIVERED ORDERS'])
                                   for r in rows if r['TRIP ID'] is not None
                                   and r['TOTAL DELIVERED ORDERS']]

    wb, stats = build_shipment_workbook(rng, shipment_trips)
    wb.save(os.path.join(HERE, 'shipments_week.xlsx'))
    raw_p, final_p, n_trips = build_example_pair()
    print(f'Wrote worked example ({n_trips} trips): {os.path.relpath(raw_p)} -> {os.path.relpath(final_p)}')
    print(f'Wrote {count} raw files to {os.path.relpath(RAW_DIR)}')
    print(f'Wrote shipment file covering {len(shipment_trips)} trips '
          f'({stats["match"]} match, {stats["over_reported"]} over-reported, '
          f'{stats["missing"]} missing)')


if __name__ == '__main__':
    main()
